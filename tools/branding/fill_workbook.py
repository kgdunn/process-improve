"""Fill the response column of a run sheet downloaded from the browser app, as an operator would."""

import sys

import numpy as np
from openpyxl import load_workbook

source, target = sys.argv[1:3]
book = load_workbook(source)
runs = book["Runs"]
header = [cell.value for cell in runs[1]]
print("columns:", header)
col = {name: i for i, name in enumerate(header)}
rng = np.random.default_rng(7)
for row in runs.iter_rows(min_row=2):
    t, p, f = (float(row[col[name]].value) for name in ("T", "P", "F"))
    # A yield that rises with temperature, falls with pressure, and bends with flow.
    y = 62 + 0.16 * (t - 175) - 3.1 * (p - 2) + 0.35 * (f - 15) - 0.09 * (f - 15) ** 2 + rng.normal(0, 0.6)
    row[col["Yield"]].value = round(y, 1)
book.save(target)
print("filled", runs.max_row - 1, "runs")
